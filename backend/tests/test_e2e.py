"""End-to-end test: upload a repo zip + PDF + DOCX, watch the live events, query the graph, chat with it."""
from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
import time
import zipfile
from pathlib import Path

import pytest

_TMP = tempfile.mkdtemp(prefix="nexmind-test-")
os.environ["NEXMIND_DATA_DIR"] = _TMP
os.environ["NEXMIND_STREAM_DELAY"] = "0"
os.environ["NEXMIND_LLM_PROVIDER"] = "offline"

from fastapi.testclient import TestClient  # noqa: E402

from nexmind.main import app  # noqa: E402

SAMPLE = Path(__file__).resolve().parent.parent / "samples" / "sample_project"


def make_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in SAMPLE.rglob("*"):
            if p.is_file():
                z.write(p, "shop-demo/" + p.relative_to(SAMPLE).as_posix())
        z.writestr("shop-demo/node_modules/junk/index.js", "function junk(){}")  # must be ignored
        z.writestr("../evil.txt", "zip-slip attempt")  # must be ignored
    return buf.getvalue()


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    lines = [
        "Payment Gateway Integration Guide",
        "The payment gateway processes card transactions for the shop. Each charge is recorded in the audit log.",
        "Stripe handles card tokenization. PCI DSS compliance requires that card data is never stored.",
        "The Database stores order history. PostgreSQL is recommended for production deployments.",
    ]
    y = 800
    for ln in lines:
        c.drawString(50, y, ln)
        y -= 24
    c.showPage()
    c.drawString(50, 800, "Refund Policy")
    c.drawString(50, 776, "Refunds are issued through the payment gateway within five business days.")
    c.drawString(50, 752, "Stripe supports partial refunds. Every refund is written to the audit log.")
    c.save()
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Operations Handbook", 1)
    d.add_paragraph("The operations team monitors the payment gateway and the audit log every morning.")
    d.add_heading("Incident response", 2)
    d.add_paragraph("If Stripe is unavailable the team switches to manual card processing and notifies security.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c
    shutil.rmtree(_TMP, ignore_errors=True)


def wait_job(client, job_id, timeout=120):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("done", "failed"):
            return j
        time.sleep(0.2)
    raise AssertionError("job timeout")


@pytest.fixture(scope="module")
def ws(client):
    r = client.post("/api/workspaces", json={"name": "Shop demo", "description": "e2e"})
    assert r.status_code == 201
    ws_id = r.json()["id"]
    with client.websocket_connect(f"/ws/workspaces/{ws_id}") as sock:
        hello = sock.receive_json()
        assert hello["type"] == "hello"
        files = [
            ("files", ("shop-demo.zip", make_zip(), "application/zip")),
            ("files", ("gateway.pdf", make_pdf(), "application/pdf")),
            ("files", ("handbook.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ]
        r = client.post(f"/api/workspaces/{ws_id}/ingest", files=files)
        assert r.status_code == 202, r.text
        job_id = r.json()["job_id"]
        events = []
        while True:
            ev = sock.receive_json()
            events.append(ev)
            if ev["type"] in ("job_done", "job_failed"):
                break
    types = [e["type"] for e in events]
    assert "job_failed" not in types, events[-1]
    assert "progress" in types and "graph_stream_start" in types and "graph_batch" in types and "graph_ready" in types
    streamed_nodes = sum(len(e["nodes"]) for e in events if e["type"] == "graph_batch")
    job = wait_job(client, job_id)
    assert job["status"] == "done", job
    assert streamed_nodes == job["result"]["nodes"]
    return {"id": ws_id, "stats": job["result"], "events": events}


def test_stats(ws):
    s = ws["stats"]
    assert s["files"] >= 8 and s["documents"] >= 4 and s["pages"] == 2
    assert s["node_types"]["Class"] >= 3
    assert s["edge_types"].get("CALLS", 0) >= 8
    assert s["edge_types"].get("IMPORTS", 0) >= 5
    assert s["edge_types"].get("INHERITS", 0) == 1
    assert s["edge_types"].get("REFERENCES", 0) >= 3  # docs -> code links
    assert s["clusters"] >= 2


def test_security_ignores(client, ws):
    g = client.get(f"/api/workspaces/{ws['id']}/graph").json()
    labels = {n["label"] for n in g["nodes"]}
    assert "junk" not in labels and "evil.txt" not in labels


def test_graph_payload(client, ws):
    g = client.get(f"/api/workspaces/{ws['id']}/graph").json()
    assert g["nodes"] and g["links"] and g["clusters"]
    n = g["nodes"][0]
    assert {"id", "type", "label", "x", "y", "z", "cluster", "score"} <= set(n)
    ids = {n["id"] for n in g["nodes"]}
    assert all(l["source"] in ids and l["target"] in ids for l in g["links"])
    # filtering
    only_code = client.get(f"/api/workspaces/{ws['id']}/graph", params={"types": "Function,Method,Class,File"}).json()
    assert {n["type"] for n in only_code["nodes"]} <= {"Function", "Method", "Class", "File"}


def _find(client, ws, q, typ=None):
    r = client.get(f"/api/workspaces/{ws['id']}/search", params={"q": q}).json()
    for n in r["nodes"]:
        if n["label"].lower() == q.lower() and (typ is None or n["type"] == typ):
            return n
    raise AssertionError(f"{q} not found: {[n['label'] for n in r['nodes']]}")


def test_node_detail_and_edges(client, ws):
    n = _find(client, ws, "place_order", "Function")
    d = client.get(f"/api/workspaces/{ws['id']}/nodes/{n['id']}").json()
    assert d["code"]["text"].startswith("def place_order")
    assert "CALLS" in d["relations"]
    callees = {r["node"]["label"] for r in d["relations"]["CALLS"]["out"]}
    assert {"validate_token", "charge_card", "Order"} <= callees
    assert any(r["node"]["label"] == "order_handler" for r in d["relations"]["CALLS"]["in"])
    assert d["breadcrumb"]
    brief = client.get(f"/api/workspaces/{ws['id']}/nodes/{n['id']}", params={"brief": "true"}).json()
    assert "relations" not in brief and brief["summary"]
    edge_id = d["relations"]["CALLS"]["out"][0]["edge_id"]
    e = client.get(f"/api/workspaces/{ws['id']}/edges/{edge_id}").json()
    assert "calls" in e["explanation"] and e["evidence"]


def test_explain_node(client, ws):
    n = _find(client, ws, "place_order", "Function")
    res = client.get(f"/api/workspaces/{ws['id']}/nodes/{n['id']}/explain")
    assert res.status_code == 200
    data = res.json()
    assert "explanation" in data and len(data["explanation"]) > 0
    assert data["provider"] in ("offline", "mock", "openai", "anthropic")


def test_doc_code_linking(client, ws):
    n = _find(client, ws, "validate_token", "Function")
    d = client.get(f"/api/workspaces/{ws['id']}/nodes/{n['id']}").json()
    refs = d["relations"].get("REFERENCES", {}).get("in", [])
    assert refs, "README section should reference validate_token"
    assert any(r["node"]["type"] == "Section" for r in refs)


def test_inheritance_and_method_resolution(client, ws):
    n = _find(client, ws, "PriorityOrder", "Class")
    d = client.get(f"/api/workspaces/{ws['id']}/nodes/{n['id']}").json()
    assert d["relations"]["INHERITS"]["out"][0]["node"]["label"] == "Order"


def test_impact(client, ws):
    n = _find(client, ws, "validate_token", "Function")
    r = client.get(f"/api/workspaces/{ws['id']}/impact", params={"node": n["id"], "direction": "upstream", "depth": 4}).json()
    labels = {x["label"] for lvl in r["levels"] for x in lvl["nodes"]}
    assert {"place_order", "order_handler"} <= labels
    assert r["risk"]["level"] in ("LOW", "MEDIUM", "HIGH", "CRITICAL")
    assert r["highlight"]["nodes"] and r["highlight"]["edges"]


def test_trace(client, ws):
    a = _find(client, ws, "order_handler", "Function")
    b = _find(client, ws, "audit", "Function")
    r = client.get(f"/api/workspaces/{ws['id']}/trace", params={"source": a["id"], "target": b["id"]}).json()
    assert r["found"] and r["nodes"][0]["label"] == "order_handler" and r["nodes"][-1]["label"] == "audit"


def test_processes_clusters_insights(client, ws):
    procs = client.get(f"/api/workspaces/{ws['id']}/processes").json()
    assert procs and procs[0]["steps"]
    cl = client.get(f"/api/workspaces/{ws['id']}/clusters").json()
    assert cl and all("center" in c for c in cl)
    ins = client.get(f"/api/workspaces/{ws['id']}/insights").json()
    assert isinstance(ins, list) and ins


def test_pdf_docx_entities(client, ws):
    g = client.get(f"/api/workspaces/{ws['id']}/graph", params={"types": "Document,Section,Concept,Entity"}).json()
    kinds = {n["type"] for n in g["nodes"]}
    assert {"Document", "Section", "Concept"} <= kinds
    titles = {n["label"] for n in g["nodes"] if n["type"] == "Document"}
    assert any("Handbook" in t or "Operations" in t for t in titles)


def test_search(client, ws):
    r = client.get(f"/api/workspaces/{ws['id']}/search", params={"q": "refund policy stripe"}).json()
    assert r["passages"] and "refund" in r["passages"][0]["snippet"].lower()


def test_chat(client, ws):
    t = client.post(f"/api/workspaces/{ws['id']}/threads", json={"title": "New chat"}).json()
    # streaming
    with client.stream("POST", f"/api/workspaces/{ws['id']}/threads/{t['id']}/messages", json={"content": "How does login and token validation work?"}) as r:
        assert r.status_code == 200
        body = "".join(r.iter_text())
    events = [blk.split("\n")[0].removeprefix("event: ") for blk in body.strip().split("\n\n")]
    assert events[0] == "start" and "retrieval" in events and "token" in events and events[-1] == "done"
    # non-streaming + impact intent
    r = client.post(f"/api/workspaces/{ws['id']}/threads/{t['id']}/messages", json={"content": "What would break if I change `validate_token`?", "stream": False}).json()
    assert r["retrieval"]["intent"] == "impact"
    assert any(a["role"] == "impact" for a in r["retrieval"]["activations"])
    assert "place_order" in r["done"]["answer"]
    # path intent
    r = client.post(f"/api/workspaces/{ws['id']}/threads/{t['id']}/messages", json={"content": "What is the path between `order_handler` and `audit`?", "stream": False}).json()
    assert r["retrieval"]["intent"] == "path"
    msgs = client.get(f"/api/workspaces/{ws['id']}/threads/{t['id']}/messages").json()
    assert len(msgs) == 6 and msgs[1]["meta"]["citations"] is not None
    assert client.get(f"/api/workspaces/{ws['id']}/threads").json()[0]["title"].startswith("How does login")


def test_exports(client, ws):
    for fmt, needle in (("json", '"nodes"'), ("graphml", "<graphml"), ("cypher", "MERGE (n:Node")):
        r = client.get(f"/api/workspaces/{ws['id']}/export", params={"format": fmt})
        assert r.status_code == 200 and needle in r.text


def test_git_url_validation(client, ws):
    for bad in ("file:///etc/passwd", "http://github.com/a/b", "https://evil.example.com/a/b", "git@github.com:a/b.git"):
        r = client.post(f"/api/workspaces/{ws['id']}/ingest/git", json={"url": bad})
        assert r.status_code == 400, bad


def test_delete_source_rebuilds(client, ws):
    srcs = client.get(f"/api/workspaces/{ws['id']}/sources").json()
    pdf_src = next(s for s in srcs if s["kind"] == "files")
    r = client.delete(f"/api/workspaces/{ws['id']}/sources/{pdf_src['id']}")
    job = wait_job(client, r.json()["job_id"])
    assert job["status"] == "done"
    assert job["result"]["documents"] == 2  # only the README + security docs inside the zip remain


# ----------------------------------------------------------------------------- v1.1: local-server onboarding features

def test_health_identifies_app(client):
    h = client.get("/api/health").json()
    assert h["app"] == "nexmind" and h["capabilities"]["local_paths"] is True and h["version"]


def test_private_network_preflight(client):
    r = client.options(
        "/api/health",
        headers={
            "Origin": "https://nexmind.vercel.app", "Access-Control-Request-Method": "GET",
            "Access-Control-Request-Private-Network": "true",
        },
    )
    assert r.status_code == 200
    assert r.headers.get("access-control-allow-private-network") == "true"
    assert r.headers.get("access-control-allow-origin") in ("*", "https://nexmind.vercel.app")


def test_local_folder_ingest_and_browser(client):
    ws_id = client.post("/api/workspaces", json={"name": "local"}).json()["id"]
    fs = client.get("/api/fs/list", params={"path": str(SAMPLE.parent)}).json()
    assert any(e["name"] == "sample_project" for e in fs["entries"]) and fs["roots"]
    bad = client.post(f"/api/workspaces/{ws_id}/ingest/local", json={"path": "/definitely/not/here"})
    assert bad.status_code == 400
    root = client.post(f"/api/workspaces/{ws_id}/ingest/local", json={"path": "/"})
    assert root.status_code == 400
    r = client.post(f"/api/workspaces/{ws_id}/ingest/local", json={"path": str(SAMPLE)})
    assert r.status_code == 202
    job = wait_job(client, r.json()["job_id"])
    assert job["status"] == "done" and job["result"]["files"] >= 8
    # deleting a local source must never touch the user's folder
    src = client.get(f"/api/workspaces/{ws_id}/sources").json()[0]
    assert src["kind"] == "local"
    wait_job(client, client.delete(f"/api/workspaces/{ws_id}/sources/{src['id']}").json()["job_id"])
    assert (SAMPLE / "shop" / "auth.py").exists()


def test_git_token_never_leaks():
    from nexmind.ingest.files import IngestError, git_clone

    with pytest.raises(IngestError) as e:
        git_clone("https://github.com/does-not-exist-zz/nope-nope", Path(_TMP) / "x", None, "ghp_SECRETSECRET1234")
    assert "SECRET" not in str(e.value)
    with pytest.raises(IngestError):
        git_clone("https://github.com/a/b", Path(_TMP) / "y", None, "bad token with spaces")


def test_docs_workspace_rejects_git(client):
    r = client.post("/api/workspaces", json={"name": "Docs Test", "kind": "docs"})
    assert r.status_code == 201
    ws_id = r.json()["id"]
    assert r.json()["kind"] == "docs"
    r2 = client.post(f"/api/workspaces/{ws_id}/ingest/git", json={"url": "https://github.com/foo/bar"})
    assert r2.status_code == 400
    assert "Document workspaces take" in r2.json()["detail"]


def test_docs_workspace_ignores_code_files(client):
    r = client.post("/api/workspaces", json={"name": "Docs Study", "kind": "docs"})
    assert r.status_code == 201
    ws_id = r.json()["id"]
    assert r.json()["kind"] == "docs"

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("notes/guide.md", "# Guide\n\nThis is a user manual for the documentation.")
        z.writestr("notes/code.py", "class Helper:\n    def run():\n        pass\n")
    zip_bytes = buf.getvalue()

    r2 = client.post(
        f"/api/workspaces/{ws_id}/ingest",
        files=[("files", ("mixed.zip", zip_bytes, "application/zip"))],
    )
    assert r2.status_code == 202
    job = wait_job(client, r2.json()["job_id"])
    assert job["status"] == "done"

    g = client.get(f"/api/workspaces/{ws_id}/graph").json()
    types = [n["type"] for n in g["nodes"]]
    assert "Document" in types
    assert "Function" not in types
    assert "Class" not in types


def test_explain_and_relayout(client, ws):
    g = client.get(f"/api/workspaces/{ws['id']}/graph").json()
    assert len(g["nodes"]) > 0
    node = g["nodes"][0]
    assert "degree" in node
    assert "betweenness" in node

    exp = client.get(f"/api/workspaces/{ws['id']}/nodes/{node['id']}/explain").json()
    assert "explanation" in exp
    assert len(exp["explanation"]) > 20

    relayout_res = client.post(f"/api/workspaces/{ws['id']}/relayout").json()
    assert relayout_res["nodes"] == len(g["nodes"])


def test_live_workspace_and_connect(client, monkeypatch):
    # 1. kind 'live' accepted and rejects git / upload
    r = client.post("/api/workspaces", json={"name": "Live Test WS", "kind": "live"})
    assert r.status_code == 201
    live_ws_id = r.json()["id"]
    assert r.json()["kind"] == "live"

    # git into live -> 400 with string detail
    r_git = client.post(f"/api/workspaces/{live_ws_id}/ingest/git", json={"url": "https://github.com/foo/bar"})
    assert r_git.status_code == 400
    assert "Live workspaces are connected from a folder." in r_git.json()["detail"]
    assert isinstance(r_git.json()["detail"], str)

    # upload into live -> 400 with string detail
    r_up = client.post(f"/api/workspaces/{live_ws_id}/ingest", files=[("files", ("test.zip", b"fake", "application/zip"))])
    assert r_up.status_code == 400
    assert "Live workspaces are connected from a folder." in r_up.json()["detail"]
    assert isinstance(r_up.json()["detail"], str)

    # 2. /api/fs/list with files=1 returns dirs then files and flags sample_project as is_project
    r_fs = client.get(f"/api/fs/list?path={SAMPLE.parent}&files=1")
    assert r_fs.status_code == 200
    fs_data = r_fs.json()
    entries = fs_data["entries"]
    assert len(entries) > 0
    # verify dirs first
    seen_file = False
    for e in entries:
        if e["kind"] == "file":
            seen_file = True
        elif e["kind"] == "dir":
            assert not seen_file, "Directories must come before files in fs/list"
    sample_entry = next((e for e in entries if e["name"] == SAMPLE.name), None)
    assert sample_entry is not None
    assert sample_entry["is_project"] is True
    assert sample_entry["markers"] is not None and len(sample_entry["markers"]) > 0

    # 3. /api/fs/check returns supported_files >= 8 for sample project and friendly 400/ok=false for "/" and missing path
    r_check = client.get(f"/api/fs/check?path={SAMPLE}")
    assert r_check.status_code == 200
    check_data = r_check.json()
    assert check_data["ok"] is True
    assert check_data["supported_files"] >= 8
    assert check_data["has_code"] is True

    r_root = client.get("/api/fs/check?path=/")
    assert r_root.status_code == 400 or r_root.json().get("ok") is False
    assert "drive root" in str(r_root.json()).lower() or "not found" in str(r_root.json()).lower()

    r_missing = client.get("/api/fs/check?path=/nonexistent_directory_123456")
    assert r_missing.status_code == 400 or r_missing.json().get("ok") is False

    # 4. /api/connect creates and reuses
    r_conn1 = client.post("/api/connect", json={"path": str(SAMPLE), "name": "Live Sample Project"})
    assert r_conn1.status_code == 200
    c1 = r_conn1.json()
    assert c1["reused"] is False
    assert c1["name"] == "Live Sample Project"
    ws1_id = c1["workspace_id"]
    wait_job(client, c1["job_id"])

    # verify kind 'live' behaves like code
    g = client.get(f"/api/workspaces/{ws1_id}/graph").json()
    types = {n["type"] for n in g["nodes"]}
    assert "Function" in types or "Class" in types or "File" in types

    # second call with same resolved path reuses
    r_conn2 = client.post("/api/connect", json={"path": str(SAMPLE)})
    assert r_conn2.status_code == 200
    c2 = r_conn2.json()
    assert c2["reused"] is True
    assert c2["workspace_id"] == ws1_id

    # 5. /api/connect/sessions returns the new session
    r_sess = client.get("/api/connect/sessions")
    assert r_sess.status_code == 200
    sess_data = r_sess.json()
    assert "now" in sess_data
    sessions = sess_data["sessions"]
    matched = [s for s in sessions if s["workspace_id"] == ws1_id]
    assert len(matched) >= 1
    assert matched[-1]["path"] == str(SAMPLE.resolve())

    # 6. watcher triggers a second rebuild after a file is added (use NEXMIND_WATCH_INTERVAL=0.2)
    from nexmind.watcher import watcher
    monkeypatch.setenv("NEXMIND_WATCH_INTERVAL", "0.2")
    monkeypatch.setenv("NEXMIND_WATCH_SETTLE", "0.4")
    watcher.wake()

    # Create a small temp watched workspace
    watch_dir = Path(tempfile.mkdtemp(prefix="nexmind-watch-test-"))
    try:
        (watch_dir / "app.py").write_text("def hello():\n    return 'world'\n", encoding="utf-8")
        r_cw = client.post("/api/connect", json={"path": str(watch_dir), "name": "Watcher Test", "watch": True})
        assert r_cw.status_code == 200
        w_info = r_cw.json()
        watch_ws_id = w_info["workspace_id"]
        wait_job(client, w_info["job_id"])

        initial_jobs = client.get(f"/api/workspaces/{watch_ws_id}/jobs").json()
        assert len(initial_jobs) >= 1

        # Add a new file to trigger watcher
        time.sleep(0.3)
        (watch_dir / "second.py").write_text("def added():\n    return 42\n", encoding="utf-8")

        # Wait for watcher to detect and submit rebuild job
        deadline = time.time() + 6.0
        rebuild_found = False
        while time.time() < deadline:
            jobs_now = client.get(f"/api/workspaces/{watch_ws_id}/jobs").json()
            rebuilds = [j for j in jobs_now if j["kind"] == "rebuild"]
            if rebuilds:
                rebuild_found = True
                wait_job(client, rebuilds[0]["id"])
                break
            time.sleep(0.2)
        assert rebuild_found, "Watcher did not trigger a rebuild after file addition"
    finally:
        try:
            client.delete(f"/api/workspaces/{ws1_id}")
        except Exception:
            pass
        try:
            client.delete(f"/api/workspaces/{watch_ws_id}")
        except Exception:
            pass
        shutil.rmtree(watch_dir, ignore_errors=True)




def test_connect_script_endpoint(client):
    r = client.get('/connect.py')
    assert r.status_code == 200, r.text
    ct = r.headers.get('content-type', '')
    assert 'text/plain' in ct, f'unexpected content-type: {ct}'
    script = r.text
    compile(script, '<connect.py>', 'exec')
    assert 'http://testserver' in script, 'base URL not injected into script'
    assert 'def main(' in script, 'no main() in script'
    assert 'urllib.request' in script
    assert 'import json' in script


def test_connect_script_no_third_party(client):
    r = client.get('/connect.py')
    assert r.status_code == 200
    for lib in ('requests', 'httpx', 'aiohttp', 'flask', 'fastapi', 'pydantic'):
        assert 'import ' + lib not in r.text


def test_fs_pick_endpoint(client, monkeypatch):
    import subprocess
    import types

    # 1. Successful pick returns normalised path
    def mock_run_success(*args, **kwargs):
        return types.SimpleNamespace(returncode=0, stdout="C:/Users/me/project/../project\n", stderr="")

    monkeypatch.setattr(subprocess, "run", mock_run_success)
    r = client.post("/api/fs/pick", json={"title": "Pick", "initial": "C:/Users"})
    assert r.status_code == 200
    data = r.json()
    assert data["cancelled"] is False
    assert data["path"] == os.path.normpath("C:/Users/me/project")

    # 2. Cancelled pick
    def mock_run_cancel(*args, **kwargs):
        return types.SimpleNamespace(returncode=0, stdout="\n", stderr="")

    monkeypatch.setattr(subprocess, "run", mock_run_cancel)
    r = client.post("/api/fs/pick", json={})
    assert r.status_code == 200
    data = r.json()
    assert data["cancelled"] is True
    assert data["path"] is None

    # 3. Tkinter fails -> 501
    def mock_run_fail(*args, **kwargs):
        return types.SimpleNamespace(returncode=1, stdout="", stderr="No display")

    monkeypatch.setattr(subprocess, "run", mock_run_fail)
    r = client.post("/api/fs/pick", json={})
    assert r.status_code == 501
    assert "not available on this computer" in r.text

    # 4. Reject non-loopback clients
    remote_client = TestClient(app, client=("192.168.1.100", 54321))
    r = remote_client.post("/api/fs/pick", json={})
    assert r.status_code == 403


def test_fs_check_src_subfolder(client, tmp_path):
    proj_dir = tmp_path / "my_project"
    proj_dir.mkdir()
    (proj_dir / "package.json").write_text('{"name": "demo"}')
    src_dir = proj_dir / "src"
    src_dir.mkdir()
    (src_dir / "index.ts").write_text('console.log("hello");')
    r = client.get(f"/api/fs/check?path={src_dir}")
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert data["exists"] is True
    assert data["is_dir"] is True
    assert data["supported_files"] > 0
    assert data["has_code"] is True
    assert "package.json" in data["markers"]


def test_ping_endpoint_and_preflight(client, monkeypatch):
    from nexmind.config import settings

    # 1. Ping OPTIONS preflight returns 204 with PNA and CORS headers
    r_options = client.options(
        "/api/ping",
        headers={
            "Origin": "https://random-site.com",
            "Access-Control-Request-Method": "GET",
            "Access-Control-Request-Private-Network": "true",
        },
    )
    assert r_options.status_code in (200, 204)
    assert r_options.headers.get("access-control-allow-origin") == "*"
    assert r_options.headers.get("access-control-allow-private-network") == "true"
    assert "GET" in r_options.headers.get("access-control-allow-methods", "")

    # 2. Ping GET from allowed origin
    monkeypatch.setattr(settings, "cors_origins", ["https://allowed.vercel.app"])
    monkeypatch.setattr(settings, "cors_origin_regex", None)
    r_allowed = client.get("/api/ping", headers={"Origin": "https://allowed.vercel.app"})
    assert r_allowed.status_code == 200
    data_allowed = r_allowed.json()
    assert data_allowed["app"] == "nexmind"
    assert data_allowed["ok"] is True
    assert data_allowed["origin_allowed"] is True
    assert data_allowed["origin"] == "https://allowed.vercel.app"
    assert r_allowed.headers.get("access-control-allow-origin") == "*"
    assert r_allowed.headers.get("access-control-allow-private-network") == "true"

    # 3. Ping GET from disallowed origin
    r_disallowed = client.get("/api/ping", headers={"Origin": "https://evil.example.com"})
    assert r_disallowed.status_code == 200
    data_disallowed = r_disallowed.json()
    assert data_disallowed["app"] == "nexmind"
    assert data_disallowed["ok"] is True
    assert data_disallowed["origin_allowed"] is False
    assert data_disallowed["origin"] == "https://evil.example.com"
    # Still carries Access-Control-Allow-Origin: * and PNA
    assert r_disallowed.headers.get("access-control-allow-origin") == "*"
    assert r_disallowed.headers.get("access-control-allow-private-network") == "true"

    # 4. Regex origin matching
    monkeypatch.setattr(settings, "cors_origin_regex", r"^https://.*\.vercel\.app$")
    r_regex = client.get("/api/ping", headers={"Origin": "https://my-preview-123.vercel.app"})
    assert r_regex.json()["origin_allowed"] is True

    # 5. Direct request without Origin header
    r_no_origin = client.get("/api/ping")
    assert r_no_origin.status_code == 200
    assert r_no_origin.json()["origin_allowed"] is True
    assert r_no_origin.json()["origin"] is None

    # 6. /api/health behavior is unchanged
    r_health = client.get("/api/health")
    assert r_health.status_code == 200
    assert r_health.json()["status"] == "ok"
