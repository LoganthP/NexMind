"""Document extraction: PDF, DOCX, Markdown, HTML, plain text, CSV/JSON, Jupyter notebooks -> ordered sections."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from ..config import settings


@dataclass
class Section:
    title: str
    text: str
    page: int | None = None
    level: int = 1


@dataclass
class DocParse:
    path: str
    doc_type: str
    title: str
    sections: list[Section] = field(default_factory=list)
    pages: int | None = None
    error: str = ""

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


def _read_text(p: Path, limit: int = 2_000_000) -> str:
    raw = p.read_bytes()[:limit]
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _clean(t: str) -> str:
    t = t.replace("\x00", "")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _group_paragraphs(text: str, max_chars: int = 2200) -> list[Section]:
    """Split headingless text into sections of paragraphs, titled by their first line."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    sections: list[Section] = []
    buf: list[str] = []
    size = 0
    for p in paras:
        if size + len(p) > max_chars and buf:
            sections.append(Section(title=_title_of(buf[0]), text="\n\n".join(buf)))
            buf, size = [], 0
        buf.append(p)
        size += len(p)
    if buf:
        sections.append(Section(title=_title_of(buf[0]), text="\n\n".join(buf)))
    return sections


def _title_of(s: str) -> str:
    first = s.strip().splitlines()[0] if s.strip() else "Section"
    first = re.sub(r"^[#>*\-\s]+", "", first)
    return (first[:70] + "…") if len(first) > 70 else first or "Section"


# ----------------------------------------------------------------------------- formats

def parse_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    cur_title, cur_level, buf = None, 1, []
    in_code = False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else re.match(r"^(#{1,6})\s+(.*\S)\s*$", line)
        if m:
            if cur_title is not None or buf:
                body = _clean("\n".join(buf))
                if body or cur_title:
                    sections.append(Section(title=cur_title or "Introduction", text=body or (cur_title or ""), level=cur_level))
            cur_title, cur_level, buf = m.group(2).strip(), len(m.group(1)), []
        else:
            buf.append(line)
    body = _clean("\n".join(buf))
    if body or cur_title:
        sections.append(Section(title=cur_title or "Introduction", text=body or (cur_title or ""), level=cur_level))
    return [s for s in sections if len(s.text.strip()) > 0]


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.sections: list[Section] = []
        self._title: str | None = None
        self._buf: list[str] = []
        self._skip = 0
        self._in_heading = False
        self._heading_buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "svg"):
            self._skip += 1
        elif tag in ("h1", "h2", "h3"):
            self._flush()
            self._in_heading = True
            self._heading_buf = []
        elif tag in ("p", "br", "li", "tr", "div"):
            self._buf.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg"):
            self._skip = max(0, self._skip - 1)
        elif tag in ("h1", "h2", "h3") and self._in_heading:
            self._in_heading = False
            self._title = " ".join("".join(self._heading_buf).split())

    def handle_data(self, data):
        if self._skip:
            return
        (self._heading_buf if self._in_heading else self._buf).append(data)

    def _flush(self):
        body = _clean("".join(self._buf))
        if body:
            self.sections.append(Section(title=self._title or "Content", text=body))
        self._buf = []
        self._title = None

    def close(self):
        super().close()
        self._flush()


def parse_html(text: str) -> list[Section]:
    p = _HTMLText()
    p.feed(text)
    p.close()
    return p.sections


def parse_pdf(path: Path) -> tuple[list[Section], int, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = len(reader.pages)
    title = ""
    try:
        title = (reader.metadata.title or "").strip() if reader.metadata else ""  # type: ignore[union-attr]
    except Exception:
        pass
    sections: list[Section] = []
    # group consecutive pages into sections of ~2500 chars so graphs stay readable on long PDFs
    buf: list[str] = []
    first_page = 1
    for i, page in enumerate(reader.pages, start=1):
        try:
            t = _clean(page.extract_text() or "")
        except Exception:
            t = ""
        if not t:
            continue
        if not buf:
            first_page = i
        buf.append(t)
        if sum(len(x) for x in buf) >= 2500:
            sections.append(Section(title=_pdf_title(buf[0], first_page, i), text="\n\n".join(buf), page=first_page))
            buf = []
    if buf:
        sections.append(Section(title=_pdf_title(buf[0], first_page, pages), text="\n\n".join(buf), page=first_page))
    return sections, pages, title


def _pdf_title(first_text: str, p1: int, p2: int) -> str:
    head = _title_of(first_text)
    pr = f"p.{p1}" if p1 == p2 else f"p.{p1}-{p2}"
    return f"{head} ({pr})"


def parse_docx(path: Path) -> list[Section]:
    import docx  # python-docx

    d = docx.Document(str(path))
    sections: list[Section] = []
    title, level, buf = None, 1, []

    def flush():
        body = _clean("\n".join(buf))
        if body:
            sections.append(Section(title=title or "Introduction", text=body, level=level))

    for para in d.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        txt = para.text.strip()
        if not txt:
            continue
        if style.startswith("heading") or style == "title":
            flush()
            buf = []
            m = re.search(r"(\d+)", style)
            title, level = txt, int(m.group(1)) if m else 1
        else:
            buf.append(txt)
    flush()
    for t in d.tables:  # tables become their own section
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        body = _clean("\n".join(rows))
        if body:
            sections.append(Section(title="Table", text=body))
    return sections


def parse_notebook(text: str) -> list[Section]:
    nb = json.loads(text)
    sections: list[Section] = []
    for cell in nb.get("cells", []):
        src = "".join(cell.get("source", []))
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            sections.extend(parse_markdown(src) or [Section(title=_title_of(src), text=src)])
        else:
            sections.append(Section(title="Code cell", text=src))
    return sections


# ----------------------------------------------------------------------------- entry

def parse_document(rel: str, path: Path, doc_type: str) -> DocParse:
    dp = DocParse(path=rel, doc_type=doc_type, title=Path(rel).stem.replace("_", " ").replace("-", " "))
    try:
        if doc_type == "pdf":
            sections, pages, title = parse_pdf(path)
            dp.sections, dp.pages = sections, pages
            bogus = re.search(r"microsoft word|\.docx?|\.indd|untitled|\.tex|\.pdf", title, re.I) if title else True
            if title and not bogus:
                dp.title = title
            elif sections:  # fall back to the first line of the first page when it looks like a heading
                first = sections[0].text.strip().splitlines()[0].strip() if sections[0].text.strip() else ""
                if 3 <= len(first) <= 90:
                    dp.title = first
        elif doc_type == "docx":
            dp.sections = parse_docx(path)
            if dp.sections and dp.sections[0].title not in ("Introduction", "Table") and dp.sections[0].level == 1:
                dp.title = dp.sections[0].title
        elif doc_type == "markdown":
            text = _read_text(path)
            dp.sections = parse_markdown(text) or _group_paragraphs(text)
            h1 = re.search(r"^#\s+(.+)$", text, re.M)
            if h1:
                dp.title = h1.group(1).strip()
        elif doc_type == "html":
            dp.sections = parse_html(_read_text(path))
        elif doc_type == "notebook":
            dp.sections = parse_notebook(_read_text(path))
        elif doc_type == "json":
            text = _read_text(path, 300_000)
            try:
                text = json.dumps(json.loads(text), indent=1)[:200_000]
            except Exception:
                pass
            dp.sections = _group_paragraphs(text, 3000)
        else:  # text / csv / yaml ...
            text = _read_text(path, 600_000)
            dp.sections = _group_paragraphs(_clean(text))
    except Exception as e:  # corrupt / encrypted files should not kill the whole ingest
        dp.error = f"{type(e).__name__}: {e}"
    cap = settings.max_sections_per_doc
    if len(dp.sections) > cap:  # merge the tail so huge documents stay navigable
        head, tail = dp.sections[: cap - 1], dp.sections[cap - 1 :]
        head.append(Section(title="Remaining content", text="\n\n".join(s.text for s in tail)))
        dp.sections = head
    return dp
